# -*- coding: utf-8 -*-
"""
產生 104技能擷取程式_修改邏輯與會議審查.docx
依據研究會議審查需求，完整整理修改緣由、處理流程、資料規格、詞庫對照、驗證結果與待決事項。
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_shading(cell, color_hex):
    """設定儲存格背景底色"""
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag.endswith('shd'):
            tcPr.remove(child)
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>'))

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    """設定儲存格內邊距 (dxa)"""
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag.endswith('tcMar'):
            tcPr.remove(child)
    tcMar_xml = f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>'
    tcPr.append(parse_xml(tcMar_xml))

def set_table_borders(table, color="D3D3D3", sz="4", val="single"):
    """設定全表格精緻淡灰色邊框"""
    tblPr = table._tbl.tblPr
    borders_xml = f'''
    <w:tblBorders {nsdecls("w")}>
        <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:left w:val="none"/>
        <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:right w:val="none"/>
        <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:insideV w:val="none"/>
    </w:tblBorders>
    '''
    tblPr.append(parse_xml(borders_xml))

def apply_row_protection(row, is_header=False):
    """防止表格跨頁斷裂與設置表頭跨頁重複"""
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
    if is_header:
        trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

def format_cell_text(cell, text, bold=False, color="000000", size=9.5, align=WD_ALIGN_PARAGRAPH.LEFT):
    """填入儲存格文字並格式化"""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.bold = bold
    run.font.name = 'Microsoft JhengHei'
    run.font.size = Pt(size)
    rId = run._r.get_or_add_rPr()
    rId.append(parse_xml(f'<w:rFonts {nsdecls("w")} w:eastAsia="Microsoft JhengHei"/>'))
    if color != "000000":
        r, g, b = int(color[0:2], 16), int(color[2:4], 16), int(color[4:6], 16)
        run.font.color.rgb = RGBColor(r, g, b)

def build_review_docx(output_path="104技能擷取程式_修改邏輯與會議審查.docx"):
    doc = docx.Document()

    # 設定頁面邊界 (2.0 cm 讓表格資訊密度適中且美觀)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
        # 頁首頁尾
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("研究會議審查技術文件（草案）｜104 職缺技能擷取系統")
        hrun.font.name = "Microsoft JhengHei"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(140, 140, 140)

        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("— 僅供研究團隊內部審查與決策討論使用 —")
        frun.font.name = "Microsoft JhengHei"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(140, 140, 140)

    # 全域樣式設定
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Microsoft JhengHei'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(40, 40, 40)
    normal_style.paragraph_format.line_spacing = 1.2
    normal_style.paragraph_format.space_after = Pt(4)

    # 輔助排版函式
    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.name = 'Microsoft JhengHei'
        run.font.size = Pt(15)
        run.font.color.rgb = RGBColor(26, 54, 93) # #1A365D Deep Navy
        rId = run._r.get_or_add_rPr()
        rId.append(parse_xml(f'<w:rFonts {nsdecls("w")} w:eastAsia="Microsoft JhengHei"/>'))
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.name = 'Microsoft JhengHei'
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(43, 108, 176) # #2B6CB0 Slate Blue
        rId = run._r.get_or_add_rPr()
        rId.append(parse_xml(f'<w:rFonts {nsdecls("w")} w:eastAsia="Microsoft JhengHei"/>'))
        return p

    def add_callout(text, title="重要提示", bg_color="EDF2F7", border_color="2B6CB0"):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        cell = table.rows[0].cells[0]
        cell.width = Inches(6.9)
        set_cell_shading(cell, bg_color)
        set_cell_margins(cell, top=100, bottom=100, left=160, right=160)
        
        tcPr = cell._tc.get_or_add_tcPr()
        borders_xml = f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
        '''
        tcPr.append(parse_xml(borders_xml))
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.2
        r_title = p.add_run(f"【{title}】")
        r_title.bold = True
        r_title.font.name = 'Microsoft JhengHei'
        r_title.font.size = Pt(9.5)
        r_title.font.color.rgb = RGBColor(43, 108, 176)
        
        r_txt = p.add_run(text)
        r_txt.font.name = 'Microsoft JhengHei'
        r_txt.font.size = Pt(9.5)
        r_txt.font.color.rgb = RGBColor(45, 55, 72)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # ----------------------------------------------------
    # 文件封面 / 抬頭資訊
    # ----------------------------------------------------
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(6)
    title_p.paragraph_format.space_after = Pt(2)
    run_t = title_p.add_run("104 技能擷取系統：修改邏輯與會議審查技術文件")
    run_t.bold = True
    run_t.font.size = Pt(19)
    run_t.font.name = 'Microsoft JhengHei'
    run_t.font.color.rgb = RGBColor(26, 54, 93)

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(10)
    run_sub = sub_p.add_run("Long Format 輸出規格、9 大原始欄位完整保留、兩份詞庫對照與新版十大分類整合方案")
    run_sub.font.size = Pt(11.5)
    run_sub.font.name = 'Microsoft JhengHei'
    run_sub.font.color.rgb = RGBColor(74, 85, 104)

    # 詮釋資料盒 (Metadata Box)
    meta_tbl = doc.add_table(rows=2, cols=2)
    meta_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r in meta_tbl.rows:
        for c in r.cells:
            set_cell_shading(c, "F7FAFC")
            set_cell_margins(c, top=60, bottom=60, left=100, right=100)
    set_table_borders(meta_tbl, color="CBD5E0", sz="4")
    format_cell_text(meta_tbl.rows[0].cells[0], "文件定位：待研究會議審查與決策之技術草案", bold=True, color="2B6CB0", size=9)
    format_cell_text(meta_tbl.rows[0].cells[1], "撰寫者：Python 資料工程師兼研究助理團隊", bold=False, color="4A5568", size=9)
    format_cell_text(meta_tbl.rows[1].cells[0], "主要修改程式：104_single_file_20260930.py (原 0909版)", bold=False, color="4A5568", size=9)
    format_cell_text(meta_tbl.rows[1].cells[1], "對應詞庫：0918基底詞庫 + Chen_grouped_09版十大分類", bold=False, color="4A5568", size=9)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    add_callout(
        "本文件是一份提交給計畫主持人與研究團隊審查的「技術修改方案與決策報告」，而非已正式核准的成果報告。文中清楚區分「原始版本行為」、「目前已實作行為」、「尚未實作的提案」與「待會議決定事項」。請團隊成員重點評估資料展開粒度、研究統計定義與後續消歧規則之採用範圍。",
        title="會議審查指引", bg_color="EBF8FF", border_color="3182CE"
    )

    # ----------------------------------------------------
    # 第一章：本次修改目的與審查重點
    # ----------------------------------------------------
    add_h1("一、本次修改目的與審查重點")
    doc.add_paragraph(
        "本研究計畫旨在從巨量 104 人力銀行職缺描述中精準萃取專業技能，並銜接國家勞動市場調查與學術計量研究。為克服過往版本在「人工抽檢不易」、「多技能合併於單一儲存格不利計量統計」以及「舊版分類缺乏細緻語意邊界」等痛點，工程團隊對單檔作業腳本進行了核心架構重構。"
    )

    # 差異總表
    t1 = doc.add_table(rows=5, cols=4)
    t1.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t1)
    apply_row_protection(t1.rows[0], is_header=True)
    for r in t1.rows[1:]:
        apply_row_protection(r)

    headers1 = ["修改構面", "修改前 (原始 0909 單檔)", "修改後 (目前已實作 0930 版)", "研究目的與實作狀態"]
    col_widths1 = [Inches(1.2), Inches(1.8), Inches(2.2), Inches(1.7)]
    for i, h in enumerate(headers1):
        cell = t1.rows[0].cells[i]
        cell.width = col_widths1[i]
        set_cell_shading(cell, "1A365D")
        format_cell_text(cell, h, bold=True, color="FFFFFF", size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)

    t1_data = [
        ("主要輸出格式", "預設僅輸出 Wide Excel 寬表，技能以「｜」管道符號拼湊在單一分類儲存格內。", "以 Long format 長表為核心交付（同時輸出 Excel 與 Parquet）；每列「一職缺 × 一技能」。", "便利計量回歸模型展開；【已實作】"),
        ("原始職缺欄位", "輸出僅保留 ID/縣市/月份；職位描述、工作技能與工具欄位皆未輸出。", "完整保留 9 個原始欄位（職位描述、工作技能、工具、刊登日期等）零截斷重複呈現。", "方便研究人員直接核對原文語境與擷取結果；【已實作】"),
        ("詞庫分類位置", "無 Category_Name 欄位，僅有寬表之中文大類表頭。", "前兩欄強制固定為詞庫原始 Category_Name 與 Subcategory_Name。", "保留 Lightcast 原始階層對照性；【已實作】"),
        ("技能主分類體系", "舊版粗略九大類分類（含未嚴謹定義之 AI技能）。", "導入新版十大技能主分類（代碼 1~10 及標準中文名稱），擴充 Grouping 稽核欄位。", "契合 Word 建議版分類框架；【已實作靜態對應】")
    ]
    for row_idx, data in enumerate(t1_data, start=1):
        row = t1.rows[row_idx]
        bg = "F7FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(data):
            c = row.cells[col_idx]
            c.width = col_widths1[col_idx]
            set_cell_shading(c, bg)
            set_cell_margins(c, top=80, bottom=80, left=100, right=100)
            align = WD_ALIGN_PARAGRAPH.CENTER if col_idx in [0, 3] else WD_ALIGN_PARAGRAPH.LEFT
            bold = True if col_idx == 0 else False
            format_cell_text(c, text, bold=bold, size=8.5, align=align)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # ----------------------------------------------------
    # 第二章：修改前後的處理流程
    # ----------------------------------------------------
    add_h1("二、修改前後的處理流程")
    doc.add_paragraph(
        "為維持資料擷取的一致性與可重現性，本次修改嚴格維持既有的字串匹配與演算法核心，僅針對輸出結構、詞庫屬性映射及資料組裝進行升級。"
    )

    add_h2("1. 處理流程對比與架構拆解")
    doc.add_paragraph(
        "系統端到端執行流程如下圖所示：\n"
        "① 讀取職缺與兩份詞庫 → ② 整理詞庫並建立十大分類對照表 → ③ 建構 Aho-Corasick 自動機 → ④ 執行字串比對與既有消歧補救 → ⑤ 組合原始職缺資料與命中技能 → ⑥ 匯出 25 欄位 Long 表。"
    )

    add_callout(
        "重要澄清：在原始 104_single_file_20260909.py 腳本中，process_county() 內部即已使用 records 串列暫存一對多的技能資料（即內部 long_df）。但原版僅包含 7 個內部欄位，且在 main() 中直接被 skills_to_wide() 彙整為寬表，原始 long_df 在程式結束時被捨棄、完全未輸出。本次修改並非「無中生有建立 long 資料結構」，而是「將內部暫存結構重構擴充，躍升為正式匯出的主力產出」。",
        title="架構澄清：原版內部早已存在 long_df", bg_color="FEFCBF", border_color="B7791F"
    )

    t2 = doc.add_table(rows=6, cols=3)
    t2.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t2)
    apply_row_protection(t2.rows[0], is_header=True)
    for r in t2.rows[1:]:
        apply_row_protection(r)

    headers2 = ["處理階段", "修改前原始狀態", "本次修改狀態與理由"]
    col_widths2 = [Inches(1.5), Inches(2.7), Inches(2.7)]
    for i, h in enumerate(headers2):
        cell = t2.rows[0].cells[i]
        cell.width = col_widths2[i]
        set_cell_shading(cell, "1A365D")
        format_cell_text(cell, h, bold=True, color="FFFFFF", size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)

    t2_data = [
        ("詞庫載入與解析", "單純讀取單一基底詞庫 (0918)，僅抓取中英文名與關鍵字。", "同時讀取 0918 基底與 Chen_grouped_09 版，預先建立 Skill_ID 映射索引，保留原始大類與新十大分類。"),
        ("AC 自動機建構", "依關鍵字建構 Trie 樹與 Failure 鏈接。", "【維持原樣】演算法結構與記憶體配置未變動，確保比對效能完全一致。"),
        ("技能比對核心", "match_skills() 結合 jieba 邊界過濾；resolve_ambiguous_by_title() 補救零命中。", "【維持原樣】比對演算法、長詞優先抑制、就近文意補救等既有邏輯 100% 保留。"),
        ("資料組合 (process_county)", "僅提取 job_id、縣市、月份，與技能結果串接為 7 欄暫存表。", "【核心重構】直接抓取原始該筆 9 個職缺欄位，與命中的每個技能逐一配對組裝成字典記錄。"),
        ("資料匯出 (main)", "強制將 long_df 轉為寬表 Excel 輸出；未儲存 long_df。", "【核心重構】主要匯出具備 25 欄位的 Long Excel 與 Parquet；寬表改為 --wide 選用輸出。")
    ]
    for row_idx, data in enumerate(t2_data, start=1):
        row = t2.rows[row_idx]
        bg = "F7FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(data):
            c = row.cells[col_idx]
            c.width = col_widths2[col_idx]
            set_cell_shading(c, bg)
            set_cell_margins(c, top=80, bottom=80, left=100, right=100)
            align = WD_ALIGN_PARAGRAPH.CENTER if col_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            format_cell_text(c, text, bold=(col_idx == 0), size=8.5, align=align)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # ----------------------------------------------------
    # 第三章：long format 與原始欄位保留方式
    # ----------------------------------------------------
    add_h1("三、Long Format 與原始欄位保留方式")
    doc.add_paragraph(
        "在新版規範中，資料粒度被嚴格定義為「一筆原始職缺 × 一個擷取到的技能」。若一筆職缺命中 N 個相異技能，則於輸出表產出 N 列；這 N 列的 9 個原始職缺欄位必須完全相同，不得有任何摘要、截斷或欄位覆蓋。"
    )

    add_h2("1. 小型資料展開範例（格式示意，非實際研究結果）")
    doc.add_paragraph(
        "假設職缺編號 ID 為「4029539」的職缺，在職位描述中同時提及並命中詞庫中 3 個真實技能：\n"
        "①「員工福利」（Skill_ID: KS1237L65TCH8606B89L）\n"
        "②「心理諮詢」（Skill_ID: KS1269S6X2RMD45V5N8R）\n"
        "③「椅子按摩」（Skill_ID: KS121CS69C4QLL85G4CS）\n"
        "其展開呈現樣貌如下表所示（展示精簡代表欄位）："
    )

    t3 = doc.add_table(rows=4, cols=6)
    t3.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t3)
    apply_row_protection(t3.rows[0], is_header=True)
    for r in t3.rows[1:]:
        apply_row_protection(r)

    headers3 = ["Category_Name", "ID", "104職位名稱", "職位描述 (原始無截斷)", "SKILL_NAME_ZH", "Skill_group10"]
    col_widths3 = [Inches(1.2), Inches(0.8), Inches(1.1), Inches(1.8), Inches(1.0), Inches(1.0)]
    for i, h in enumerate(headers3):
        cell = t3.rows[0].cells[i]
        cell.width = col_widths3[i]
        set_cell_shading(cell, "1A365D")
        format_cell_text(cell, h, bold=True, color="FFFFFF", size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)

    t3_data = [
        ("Human Resources", "4029539", "人資專員", "負責規劃員工福利、提供同仁心理諮詢及預約椅子按摩紓壓...", "員工福利", "管理、商業、法遵與治理技能"),
        ("Health Care", "4029539", "人資專員", "負責規劃員工福利、提供同仁心理諮詢及預約椅子按摩紓壓...", "心理諮詢", "醫療、健康與照護技能"),
        ("Personal Care", "4029539", "人資專員", "負責規劃員工福利、提供同仁心理諮詢及預約椅子按摩紓壓...", "椅子按摩", "營運、技職與第一線服務技能")
    ]
    for row_idx, data in enumerate(t3_data, start=1):
        row = t3.rows[row_idx]
        bg = "F7FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(data):
            c = row.cells[col_idx]
            c.width = col_widths3[col_idx]
            set_cell_shading(c, bg)
            set_cell_margins(c, top=80, bottom=80, left=80, right=80)
            align = WD_ALIGN_PARAGRAPH.CENTER if col_idx in [1, 4] else WD_ALIGN_PARAGRAPH.LEFT
            format_cell_text(c, text, bold=False, size=8.0, align=align)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    add_h2("2. 原始 9 大欄位保真與來源別名正規化")
    doc.add_paragraph(
        "系統強制保留以下 9 個原始職缺欄位：ID、資料月份、刊登日期、工作角色、104職位名稱、104職位名稱碼、職位描述、工作技能、電腦工具。針對來源資料常見的欄位異質性，採用以下嚴密機制處理："
    )
    doc.add_paragraph(
        "• 欄位優先順序與別名容錯：\n"
        "  - ID：優先採用同名「ID」，缺漏時退回採用「工作編號」。若兩者均不存在，程式立即終止並拋出 ValueError，絕不隨機編造假 ID。\n"
        "  - 104職位名稱：優先採用「104職位名稱」，無此欄時退回採用「職位名稱」。\n"
        "  - 104職位名稱碼：優先採用「104職位名稱碼」，無此欄時退回採用「職位名稱碼」。\n"
        "  - 電腦工具：優先採用「電腦工具」，無此欄時退回採用「擅長工具」。\n"
        "  - 衝突檢測機制：當主要欄位與替代欄位同時存在且內容不一致時（例如 ID='A101' 且 工作編號='B202'），系統發出 WARNING 日誌記錄衝突筆數與內容，絕不靜默掩蓋。\n"
        "• 數值型別與前導零保護：以字串形態讀取 ID 與職位代碼，並透過自訂 clean 函式防護，嚴格禁止 pandas 將代碼轉為浮點數（如防止 '00201' 變為 201.0 或丟失開頭 0），且杜絕 'nan' 或 'None' 假字串。\n"
        "• 資料月份 vs 月份雙軌制：第 4 欄「資料月份」完整保留原始輸入檔內登載之年月原值；第 13 欄「月份」維持從職缺檔名（如 cleaned_彰化縣_202608.xlsx）解析出的統計月份。兩者獨立並存，絕不互相覆蓋。"
    )

    add_h2("3. 杜絕 Merge 錯配與去重範圍釐清")
    doc.add_paragraph(
        "【為什麼嚴禁使用 pd.merge？】\n"
        "若原始職缺資料中存在重複 ID（例如企業重複刊登同一職缺、或資料抓取時的重複列），若先抽取技能後再透過 pd.merge(long_df, raw_df, on='ID') 合併，會觸發 SQL 中的 Many-to-Many 交叉乘積，造成列數幾何級數膨脹，甚至將職缺 A 的技能配到職缺 B 的描述上。\n"
        "【目前程式實作方式】\n"
        "程式在 process_county() 中以逐列（iterative）方式遍歷每一筆原始職缺，取得該列的原始資料物件 raw_row 後，直接與該列命中的每個 matched 技能組合成字典並 append 至 records。完全不使用 DataFrame merge，因此即使輸入檔有 100 筆相同 ID，每筆職缺的描述與技能皆各自精準綁定，列數絕對不膨脹。\n"
        "【去重範圍技術細節】\n"
        "程式內部的去重集合 seen_ids 是在「每一筆職缺列（row）」的維度生效。換言之，同一筆職缺內不會重複出現相同的 Skill_ID；但若原始資料中原本就存在 2 筆相同 ID 的獨立職缺紀錄，兩筆職缺將各自獨立產出技能列。此處涉及後續計量母體定義，列入會議待決事項。"
    )

    # ----------------------------------------------------
    # 第四章：輸出欄位規格
    # ----------------------------------------------------
    add_h1("四、輸出欄位規格（25 欄固定字典表）")
    doc.add_paragraph(
        "目前已實作的 104_single_file_20260930.py 產出的 Long format（Excel 與 Parquet）均嚴格鎖定為以下 25 個欄位，其順序、來源與學術用途如下表所定義："
    )

    t4 = doc.add_table(rows=26, cols=5)
    t4.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t4)
    apply_row_protection(t4.rows[0], is_header=True)
    for r in t4.rows[1:]:
        apply_row_protection(r)

    headers4 = ["序", "欄位名稱", "來源", "資料型態", "欄位定義與研究用途說明"]
    col_widths4 = [Inches(0.4), Inches(1.8), Inches(1.0), Inches(0.9), Inches(2.8)]
    for i, h in enumerate(headers4):
        cell = t4.rows[0].cells[i]
        cell.width = col_widths4[i]
        set_cell_shading(cell, "1A365D")
        format_cell_text(cell, h, bold=True, color="FFFFFF", size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)

    t4_data = [
        ("1", "Category_Name", "詞庫原始", "string", "詞庫 Lightcast 原始英文大分類名稱（置於首欄供快速檢索）。"),
        ("2", "Subcategory_Name", "詞庫原始", "string", "詞庫 Lightcast 原始英文次分類名稱。"),
        ("3", "ID", "原始職缺", "string", "職缺識別碼（優先讀取 ID，無則讀取工作編號，保留前導零）。"),
        ("4", "資料月份", "原始職缺", "string", "輸入檔案內原始登載之資料月份（保留原值）。"),
        ("5", "刊登日期", "原始職缺", "string", "職缺原始刊登發布日期。"),
        ("6", "工作角色", "原始職缺", "string", "職缺原始所屬工作角色或職務類別。"),
        ("7", "104職位名稱", "原始職缺", "string", "正規化後之 104 職位名稱（優先讀取 104職位名稱，無則職位名稱）。"),
        ("8", "104職位名稱碼", "原始職缺", "string", "104 職位官方代碼（保留前導零字串，杜絕小數點浮點數）。"),
        ("9", "職位描述", "原始職缺", "string", "原始工作職位描述全文（完整保留換行與符號，絕不截斷）。"),
        ("10", "工作技能", "原始職缺", "string", "原始工作技能說明欄位全文（無截斷）。"),
        ("11", "電腦工具", "原始職缺", "string", "原始擅長工具/電腦工具欄位全文（無截斷）。"),
        ("12", "縣市", "檔名擷取", "string", "由職缺檔案名稱解析出的所屬縣市別（如：彰化縣）。"),
        ("13", "月份", "檔名擷取", "string", "由職缺檔案名稱解析出之標準化年月（如：202608）。"),
        ("14", "SKILL_ID", "技能比對", "string", "詞庫中技能唯一代碼（如 104_004512 或 KS 代碼）。"),
        ("15", "SKILL_NAME", "技能比對", "string", "詞庫官方英文標準技能名稱。"),
        ("16", "SKILL_NAME_ZH", "技能比對", "string", "詞庫官方繁體中文標準技能名稱。"),
        ("17", "SKILL_TYPE", "技能比對", "string", "技能類型（如 Hard Skill、Specialized Skill、Common Skill）。"),
        ("18", "Category_Code", "詞庫原始", "Int64 (可空)", "詞庫原始大分類數字代碼。"),
        ("19", "Subcategory_Code", "詞庫原始", "Int64 (可空)", "詞庫原始次分類數字代碼。"),
        ("20", "Skill_group10_code", "Word 十大類", "Int64 (可空)", "新版十大技能主分類數字代碼 (1 ~ 10)。"),
        ("21", "Skill_group10", "Word 十大類", "string", "新版十大技能主分類標準中文名稱。"),
        ("22", "Grouping_rule", "分類對照", "string", "十大分類判斷規則依據（如：詞庫已分類、未對應）。"),
        ("23", "Grouping_status", "分類對照", "string", "分類映射狀態（已對應、未對應）。"),
        ("24", "MATCHED_FROM", "比對程序", "string", "觸發命中的職缺來源欄位（職位描述、工作技能、工具欄、就近文意、職稱消歧）。"),
        ("25", "IS_SOFTWARE", "詞庫屬性", "boolean", "是否屬於軟體/資訊科技工具（True / False）。")
    ]
    for row_idx, data in enumerate(t4_data, start=1):
        row = t4.rows[row_idx]
        bg = "F7FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(data):
            c = row.cells[col_idx]
            c.width = col_widths4[col_idx]
            set_cell_shading(c, bg)
            set_cell_margins(c, top=50, bottom=50, left=60, right=60)
            align = WD_ALIGN_PARAGRAPH.CENTER if col_idx in [0, 2, 3] else WD_ALIGN_PARAGRAPH.LEFT
            format_cell_text(c, text, bold=(col_idx == 1), size=7.8, align=align)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    add_callout(
        "關鍵提醒：第 1、2 欄的 Category_Name / Subcategory_Name 是 Lightcast 原始分類體系；第 20、21 欄的 Skill_group10 則是本計畫修訂之新版十大主分類。兩者在研究統計上代表不同維度，獨立並存、絕不互相覆蓋！\n"
        "此外，第 24 欄 MATCHED_FROM 僅表示該技能是從哪一個來源欄位（或消歧規則）觸發，並不能作為該技能在文章中具體起訖字符位置（Span）的精確語意證據。",
        title="欄位解讀注意事項", bg_color="FFF5F5", border_color="E53E3E"
    )

    # ----------------------------------------------------
    # 第五章：兩份詞庫的使用與合併規則
    # ----------------------------------------------------
    add_h1("五、兩份詞庫的使用與合併規則")
    doc.add_paragraph(
        "本系統在詞庫層級採用「雙軌整合」策略：\n"
        "1. 基底詞庫 (0918 版)：詞條數 108,386 筆，涵蓋完整的 Skill_ID、英文名、中文名、觸發關鍵字 (Keywords) 以及 2 字詞放行名單，是建構 AC 自動機與字串掃描的唯一起點。\n"
        "2. 分類詞庫 (Chen_grouped_09 版)：收錄 42,208 列（涵蓋 26,320 筆相異 Skill_ID），提供了人工與半自動標註之十大主分類 metadata (Skill_group10、Skill_group10_code、Grouping_rule)。\n"
        "3. 規範指引 (Word 建議版)：提供十大分類之邊界定義（Include / Exclude）與消歧原則。"
    )

    add_h2("1. 避免列數膨脹的字典預載入機制")
    doc.add_paragraph(
        "若直接使用 pandas.merge 將基底詞庫與分類詞庫進行表層合併，由於 Chen 版詞庫存在多重關鍵字展開或同 ID 多列之情形，極易導致詞庫列數倍增，進而拖垮比對速度。\n"
        "【程式實作方案】\n"
        "程式在 load_automaton() 中，預先將 Chen 版詞庫依 Skill_ID 聚合為單一字典 dict[Skill_ID -> Metadata]。在掃描到技能命中時，直接透過 O(1) 字典查表將分類資訊填入記錄中。經實測，十大分類字典包含 26,320 個唯一 Skill_ID，且衝突率為 0，完全消除列數膨脹風險。"
    )

    add_h2("2. 5 個特定 ID 孤立性核查報告")
    doc.add_paragraph(
        "依據指示，工程團隊對以下 5 個特定的本土化醫療與社工技能 ID 進行了跨詞庫雙向比對實證："
    )

    t5 = doc.add_table(rows=6, cols=5)
    t5.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t5)
    apply_row_protection(t5.rows[0], is_header=True)
    for r in t5.rows[1:]:
        apply_row_protection(r)

    headers5 = ["Skill_ID", "0918 基底中文名稱", "原始 Category / Subcategory", "Chen 版存在狀態", "目前程式處理方式"]
    col_widths5 = [Inches(1.2), Inches(1.5), Inches(1.8), Inches(1.1), Inches(1.3)]
    for i, h in enumerate(headers5):
        cell = t5.rows[0].cells[i]
        cell.width = col_widths5[i]
        set_cell_shading(cell, "1A365D")
        format_cell_text(cell, h, bold=True, color="FFFFFF", size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)

    t5_data = [
        ("TW_TCM_002", "中醫特殊療法與科別", "Health Care / Alternative Therapy", "不存在 (無此 ID)", "保留技能，分類標記【未對應】"),
        ("TW_NUR_010", "病患移位與基礎照護技術", "Health Care / Nursing & Patient Care", "不存在 (無此 ID)", "保留技能，分類標記【未對應】"),
        ("TW_MED_002", "健保申報與門診行政", "Health Care / Health Care Administration", "不存在 (無此 ID)", "保留技能，分類標記【未對應】"),
        ("TW_SOC_002", "身心障礙職業重建服務", "Social Services / Government Assistance", "不存在 (無此 ID)", "保留技能，分類標記【未對應】"),
        ("TW_PERS_001", "中式整復推拿", "Personal Care / Beauty & Body Treatments", "不存在 (無此 ID)", "保留技能，分類標記【未對應】")
    ]
    for row_idx, data in enumerate(t5_data, start=1):
        row = t5.rows[row_idx]
        bg = "F7FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(data):
            c = row.cells[col_idx]
            c.width = col_widths5[col_idx]
            set_cell_shading(c, bg)
            set_cell_margins(c, top=60, bottom=60, left=80, right=80)
            align = WD_ALIGN_PARAGRAPH.CENTER if col_idx in [0, 3] else WD_ALIGN_PARAGRAPH.LEFT
            bold = True if col_idx == 0 else False
            format_cell_text(c, text, bold=bold, size=8.0, align=align)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    doc.add_paragraph(
        "【核實結論與處理原則】\n"
        "經查核確認，上述 5 筆技能「100% 僅存在於 0918 基底詞庫，Chen 版分類詞庫中完全無收錄」。\n"
        "目前程式遵循嚴格的研究保真原則：若職缺命中此 5 類技能，技能列完整保留、不予刪除；但其 Skill_group10_code 填入空值、Skill_group10 填入空字串、Grouping_status 標註為「未對應」。系統絕不自行武斷替其編造主分類，等待會議專家判定後補齊。"
    )

    # ----------------------------------------------------
    # 第六章：新十大分類與 Word 規則的採用範圍
    # ----------------------------------------------------
    add_h1("六、新十大分類與 Word 規則的採用範圍")
    doc.add_paragraph(
        "Word 建議文件《10大技能分類與選擇規則_建議版_完成.docx》建立了嚴謹的雙層架構：第一層為詞庫層級的 GR1~GR7 規則；第二層為職缺語境消歧的 JM1~JM6 規則。在審查中必須清楚交代目前的實作邊界。"
    )

    add_h2("1. 正式十大主分類體系與舊版差異")
    doc.add_paragraph(
        "Word 文件定義之 10 個標準主分類代碼與名稱如下表所示。必須特別強調：新十類並非舊九類的簡單名稱置換。舉例而言，舊版的「AI & Big Data Skills」已被拆解重整入第 8 類（資訊科技、軟體、資料與AI技能），而一般文書處理則嚴格劃入第 6 類（行政與一般數位技能）；兩者代碼與語意邊界皆不相同，不能直接沿用舊代碼！"
    )

    t6 = doc.add_table(rows=11, cols=3)
    t6.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t6)
    apply_row_protection(t6.rows[0], is_header=True)
    for r in t6.rows[1:]:
        apply_row_protection(r)

    headers6 = ["代碼", "十大技能主分類標準中文名稱", "核心範疇與決策邊界 (Scope & Boundary)"]
    col_widths6 = [Inches(0.6), Inches(2.5), Inches(3.8)]
    for i, h in enumerate(headers6):
        cell = t6.rows[0].cells[i]
        cell.width = col_widths6[i]
        set_cell_shading(cell, "1A365D")
        format_cell_text(cell, h, bold=True, color="FFFFFF", size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)

    t6_data = [
        ("1", "認知、分析與創新技能", "一般推理、問題解決、決策、研究、概念分析與非特定領域創新（排除特定科技或財務建模）。"),
        ("2", "溝通、語言與人際技能", "溝通、外語、翻譯實務、協商、團隊合作、教學、諮商（核心為人際理解與互動）。"),
        ("3", "個人特質、自我管理與身體能力", "工作態度、適應力、主動性、細心度、耐力與肢體搬運（衡量個人工作表現模式）。"),
        ("4", "財務、會計與經濟技能", "會計、審計、稅務、預算、銀行投資、財務風控（核心產出為資金、計量或貨幣決策）。"),
        ("5", "管理、商業、法遵與治理技能", "組織領導、專案管理、人資、商業策略、合約管理與法規遵循（統整組織權責）。"),
        ("6", "行政與一般數位技能", "文書行政、排程、微軟 Office、電子郵件、行事曆與基礎數位素養（常規辦公生產力）。"),
        ("7", "營運、技職與第一線服務技能", "機台操作、設備維修、營造技工、物流倉儲、餐旅零售、第一線服務作業。"),
        ("8", "資訊科技、軟體、資料與AI技能", "程式開發、演算法、資料庫、雲端、資安、資料科學、AI/ML（專業資通訊技術）。"),
        ("9", "醫療、健康與照護技能", "醫學、護理、藥學、臨床診斷、復健治療、病患照護（以實質醫療行為為核心）。"),
        ("10", "工程、自然科學、設計、環境與能源技能", "非資通訊之工程學、自然科學實驗、建築技術、材料、環保能源與技術設計。")
    ]
    for row_idx, data in enumerate(t6_data, start=1):
        row = t6.rows[row_idx]
        bg = "F7FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(data):
            c = row.cells[col_idx]
            c.width = col_widths6[col_idx]
            set_cell_shading(c, bg)
            set_cell_margins(c, top=50, bottom=50, left=70, right=70)
            align = WD_ALIGN_PARAGRAPH.CENTER if col_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            format_cell_text(c, text, bold=(col_idx == 0), size=8.0, align=align)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    add_h2("2. 實際採用範圍與演算法已知限制")
    doc.add_paragraph(
        "【目前已實作範圍】\n"
        "• 靜態詞庫映射：目前單檔腳本係採用 Chen_grouped_09 詞庫中已完成之標註作為元資料來源，於比對命中後靜態填入 Skill_group10。\n"
        "• 保留既有補救規則：保留原版針對籠統字（如「開發」、「操作」、「管控」）之就近文意與職稱輔助消歧。\n"
        "【尚未實作範圍（重要審查點）】\n"
        "• Word 第 10 節動態消歧體系（JM1~JM6）：Word 文件要求在抽取階段動態分析職位中類/大類、局部句子結構進行語意仲裁（例如將「電子郵件軟體」依職務動態決定分至 G6 或 G8）。目前單檔程式尚未實作此套複雜的動態職缺消歧器，不能聲稱已具備 Word 第 10 節之能力。\n"
        "• 既有去重演算法限制（seen_zh_terms）：\n"
        "  在 match_skills() 內部，程式宣告了 seen_zh_terms 集合。若同一個中文詞彙在詞庫中對應多個相異技能（例如「設計」同時對應服裝設計與軟體架構設計），由於程式先掃描到的詞彙會立即加入 seen_zh_terms，後續的同名候選會被直接 pass 略過。這項既有邏輯可能提早截斷了候選集合，阻礙了後續語境消歧的空間，列為研究團隊必須知悉之技術限制。"
    )

    # ----------------------------------------------------
    # 第七章：驗證結果與研究影響
    # ----------------------------------------------------
    add_h1("七、驗證結果與研究影響")
    doc.add_paragraph(
        "為確認修改後的腳本 104_single_file_20260930.py 達到生產品質與研究保真度，工程團隊構建了自動化測試套件（scripts/verify_single_file_upgrade.py），執行了 8 大核心測項，全部取得 PASS 通過。"
    )

    t7 = doc.add_table(rows=9, cols=5)
    t7.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t7)
    apply_row_protection(t7.rows[0], is_header=True)
    for r in t7.rows[1:]:
        apply_row_protection(r)

    headers7 = ["測試項目", "預期結果", "實際測試結果", "量化測試證據", "狀態"]
    col_widths7 = [Inches(1.4), Inches(1.5), Inches(1.5), Inches(1.8), Inches(0.7)]
    for i, h in enumerate(headers7):
        cell = t7.rows[0].cells[i]
        cell.width = col_widths7[i]
        set_cell_shading(cell, "1A365D")
        format_cell_text(cell, h, bold=True, color="FFFFFF", size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)

    t7_data = [
        ("1. 多技能拆列粒度", "1 職缺命中 3 技能恰好產出 3 列。", "精確產出 3 列，無合併儲存格。", "測試職缺命中福利/諮詢/按摩，輸出列數=3。", "PASS"),
        ("2. 原始 9 欄一致性", "展開後 3 列原始 9 欄與來源完全一致。", "文字與字元 100% 吻合，零截斷。", "測試職缺不一致=0；Demo 抽查 50 筆不一致=0。", "PASS"),
        ("3. 重複 ID 隔離測試", "來源重複 ID 不造成交叉乘積或列數膨脹。", "輸入 2 筆相同 ID，輸出恰好 3 列。", "各職缺技能各自獨立綁定，交叉配對數=0。", "PASS"),
        ("4. 雙格式鏡像一致", "Excel 與 Parquet 欄位與數值完全一致。", "25 欄名稱順序相同，數值對齊。", "3,955 列對齊，全表格差異儲存格數=0。", "PASS"),
        ("5. 欄位衝突檢測", "主要與替代欄位不一致時發出 Warning。", "成功攔截並記錄 4 處預設衝突。", "日誌完整打印衝突筆數與欄位名稱。", "PASS"),
        ("6. 缺失 ID 報錯防護", "缺失 ID 報錯；缺失非識別欄位自動補空。", "拋出 ValueError；補空正常。", "驗證無編造假 ID；補空欄位符合 25 欄。", "PASS"),
        ("7. 全檔零命中保護", "若整檔職缺均無技能，產出 25 欄空表。", "產出 0 列、25 欄位空 DataFrame。", "避免程式 crash，且維持下游 schema 穩定。", "PASS"),
        ("8. 修改前後比對等價", "字串擷取集合與原始版本 100% 一致。", "命中技能完全吻合，零誤差。", "Demo 500 筆測試集差異集合（舊-新=0, 新-舊=0）。", "PASS")
    ]
    for row_idx, data in enumerate(t7_data, start=1):
        row = t7.rows[row_idx]
        bg = "F7FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(data):
            c = row.cells[col_idx]
            c.width = col_widths7[col_idx]
            set_cell_shading(c, bg)
            set_cell_margins(c, top=50, bottom=50, left=60, right=60)
            align = WD_ALIGN_PARAGRAPH.CENTER if col_idx in [0, 4] else WD_ALIGN_PARAGRAPH.LEFT
            bold = True if col_idx == 4 else (True if col_idx == 0 else False)
            color = "276749" if col_idx == 4 else "000000"
            format_cell_text(c, text, bold=bold, color=color, size=7.8, align=align)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    add_callout(
        "測試範圍宣告：上述驗證係基於構造測試案例與 demo_彰化縣_202608_500.xlsx（500 筆職缺）之實測結果。本測試證明了資料結構與演算法在該資料集下的正確性，但不能直接等同於全國數百萬筆職缺全量資料集的跑批驗證；全量資料之效能與邊界案例仍屬【尚未驗證】範圍。",
        title="測試範圍核實說明", bg_color="EDF2F7", border_color="718096"
    )

    add_h2("研究統計重要注意事項（敬請計量研究成員詳讀）")
    doc.add_paragraph(
        "1. 列數不等於職缺數：Long 表每一列代表「一個技能的出現」，而非一筆職缺。絕對不可直接使用 Long 表的總列數作為職缺總數進行回歸或推論！\n"
        "2. 技能滲透率加總必定超過 100%：由於單一職缺通常同時要求多種不同分類的技能（例如同時要求人際溝通 G2 與資訊科技 G8），若計算各十大分類的職缺占比，各類占比加總必定會大於 100%。\n"
        "3. 覆蓋率計算之母體問題：目前的 Long 表僅包含「有命中技能」的列（本次 500 筆職缺中命中 496 筆，命中率 99.2%）。若要計算整體市場的技能覆蓋率，必須結合未命中的 4 筆職缺母體計算，不可單獨以 Long 表除算。\n"
        "4. 區分「未命中職缺」與「未對應技能」：\n"
        "   - 未命中職缺（Zero-hit Jobs）：整篇 JD 找不到任何詞庫技能，不在主要 Long 表中。\n"
        "   - 未對應技能（Unmapped Skills）：職缺有命中技能，但該技能在 Chen 版詞庫中缺少十大分類（如 TW_TCM_002），此時仍會產出一列，但其分類碼為空、狀態為「未對應」。"
    )

    # ----------------------------------------------------
    # 第八章：會議待決事項 (決策表)
    # ----------------------------------------------------
    add_h1("八、會議待決事項（決策表）")
    doc.add_paragraph(
        "為使本次技術審查能有效落實為後續研究決策，工程團隊整理出 8 項核心議題。所有「會議決議」欄位目前均保持空白或標註「待討論」，留待會議中由主持人與研究成員議決："
    )

    t8 = doc.add_table(rows=9, cols=5)
    t8.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t8)
    apply_row_protection(t8.rows[0], is_header=True)
    for r in t8.rows[1:]:
        apply_row_protection(r)

    headers8 = ["議題編號與名稱", "目前技術提案 / 實作現況", "研究與統計影響", "會議決議 (留空待填)", "後續負責人與期限"]
    col_widths8 = [Inches(1.2), Inches(1.7), Inches(1.8), Inches(1.1), Inches(1.1)]
    for i, h in enumerate(headers8):
        cell = t8.rows[0].cells[i]
        cell.width = col_widths8[i]
        set_cell_shading(cell, "1A365D")
        format_cell_text(cell, h, bold=True, color="FFFFFF", size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)

    t8_data = [
        ("D1. Long 格式與欄位順序", "固定 25 欄位，Category 居首，保留 9 大原始職缺欄位。", "統一團隊分析資料規格，提升人工檢查便利度。", "【待討論】", "資料工程組"),
        ("D2. 詞庫整合來源方案", "以 0918 為基底進行比對，結合 Chen 版進行十大分類映射。", "保留最豐富之關鍵字，同時享有最新分類成果。", "【待討論】", "詞庫研究組"),
        ("D3. 零命中職缺處置", "主要 Long 表僅存有命中列；零命中職缺不產出無效假列。", "若計量需分析無技能職缺，需另行產出未命中清單。", "【待討論】", "計量分析組"),
        ("D4. 缺少分類技能之補齊", "未收錄於 Chen 版的技能（如中醫等 5 筆）目前標註「未對應」。", "避免機器揣測分類失真；需人工專家審查歸類。", "【待討論】", "詞庫審查專家"),
        ("D5. 分類衝突與多重對應", "若同 Skill_ID 在詞庫有不同分類，目前優先採納第一筆合法紀錄。", "少數邊界技能可能存在爭議，需確認仲裁準則。", "【待討論】", "詞庫研究組"),
        ("D6. Wide 寬表保留需求", "預設不輸出 Wide，提供 --wide 參數供研究員自由啟用。", "節省伺服器 I/O 與硬碟空間，同時保留向下相容性。", "【待討論】", "研究助理組"),
        ("D7. Word 第10節消歧實作", "目前單檔程式暫不實作動態上下文消歧，列入第二階段評估。", "若提前引入複雜規則可能降低單檔執行吞吐量。", "【待討論】", "技術與計量組"),
        ("D8. 重複 ID 之母體計算定義", "目前以「單列職缺」去重；重複 ID 職缺各自獨立產出技能列。", "影響研究推估全國職缺總數時的母體去重口徑。", "【待討論】", "計畫主持人 / 老師")
    ]
    for row_idx, data in enumerate(t8_data, start=1):
        row = t8.rows[row_idx]
        bg = "F7FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(data):
            c = row.cells[col_idx]
            c.width = col_widths8[col_idx]
            set_cell_shading(c, bg)
            set_cell_margins(c, top=60, bottom=60, left=70, right=70)
            align = WD_ALIGN_PARAGRAPH.CENTER if col_idx in [0, 3] else WD_ALIGN_PARAGRAPH.LEFT
            bold = True if col_idx in [0, 3] else False
            color = "C53030" if col_idx == 3 else "000000"
            format_cell_text(c, text, bold=bold, color=color, size=7.8, align=align)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 尾頁結語
    p_end = doc.add_paragraph()
    p_end.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_end = p_end.add_run("文件製作完成日期：2026 年 09 月 30 日 ｜ 104 技能擷取研究團隊")
    r_end.font.size = Pt(9)
    r_end.font.color.rgb = RGBColor(120, 120, 120)

    # 儲存檔案
    doc.save(output_path)
    print(f"成功產出文件：{output_path}")

if __name__ == "__main__":
    build_review_docx()
