"""Lossless long-row import and explicit raw-job backfill; no skill inference."""

from contextlib import closing
from datetime import date, datetime
import csv
import json
import math
import os
from pathlib import Path
import sqlite3
import tempfile

JOB_COLUMNS = ["ID", "資料月份", "刊登日期", "工作角色", "104職位名稱",
               "104職位名稱碼", "職位描述", "工作技能", "電腦工具"]
SKILL_COLUMNS = ["category", "subcategory", "SKILL_ID", "SKILL_NAME", "Skill_Name_ZH",
                 "SKILL_TYPE", "Category_Code", "Subcategory_Code", "Skill_group10_code",
                 "Skill_group10", "Grouping_rule", "Grouping_status", "MATCHED_FROM",
                 "IS_SOFTWARE", "縣市", "月份", "MATCHED_KEYWORD", "START_POS", "END_POS"]
ALIASES = {
    "ID": ["ID", "工作編號", "id"],
    "資料月份": ["資料月份", "月份", "month"],
    "104職位名稱": ["104職位名稱", "職位名稱", "job_title"],
    "104職位名稱碼": ["104職位名稱碼", "職位名稱碼", "job_code"],
    "職位描述": ["職位描述", "job_desc"],
    "電腦工具": ["電腦工具", "擅長工具", "tools"],
    "工作技能": ["工作技能", "job_skills"],
    "category": ["category", "Category_Name", "SKILL_CATEGORY_NAME"],
    "subcategory": ["subcategory", "Subcategory_Name", "SKILL_SUBCATEGORY_NAME"],
    "Skill_Name_ZH": ["Skill_Name_ZH", "SKILL_NAME_ZH"],
    "Category_Code": ["Category_Code", "SKILL_CATEGORY"],
    "Subcategory_Code": ["Subcategory_Code", "SKILL_SUBCATEGORY"],
}


def quote(name):
    return '"' + name.replace('"', '""') + '"'


def value(row, key):
    """An existing NULL alias stays NULL; do not fall through to another alias."""
    for alias in ALIASES.get(key, [key]):
        if alias in row:
            return row[alias]
    return None


def scalar(v):
    if v is None or (isinstance(v, float) and math.isnan(v)):
        return None
    if isinstance(v, (date, datetime)):
        return v.isoformat()
    if isinstance(v, bool):
        return int(v)
    if isinstance(v, (str, int, float, bytes)):
        return v
    return str(v)


def text(v):
    v = scalar(v)
    return None if v is None else str(v)


def boolean(v):
    if v is None:
        return None
    if isinstance(v, bool):
        return int(v)
    s = str(v).strip().lower()
    if s in {"1", "true", "yes"}:
        return 1
    if s in {"0", "false", "no"}:
        return 0
    raise ValueError(f"Invalid IS_SOFTWARE value: {v!r}")


def payload(row):
    return json.dumps(row, ensure_ascii=False, sort_keys=True, default=str)


def job_values(row, month=None):
    result = [text(value(row, c)) for c in JOB_COLUMNS]
    if result[1] is None and month is not None:
        result[1] = str(month)
    if result[0] in (None, "") or result[1] in (None, ""):
        raise ValueError("ID and 資料月份 are required; use --month for a missing month column")
    return result


def columns(conn, table):
    return {r[1] for r in conn.execute(f"PRAGMA table_info({quote(table)})")}


def initialize(conn):
    conn.execute("PRAGMA foreign_keys=ON")
    job_defs = ",".join(f"{quote(c)} TEXT" for c in JOB_COLUMNS)
    skill_defs = ",".join(f"{quote(c)} {'INTEGER' if c in {'IS_SOFTWARE', 'START_POS', 'END_POS'} else 'TEXT'}"
                          for c in SKILL_COLUMNS)
    conn.executescript(f"""
        CREATE TABLE IF NOT EXISTS jobs ({job_defs}, PRIMARY KEY(ID, 資料月份));
        CREATE TABLE IF NOT EXISTS skill_records (
            record_id INTEGER PRIMARY KEY, ID TEXT NOT NULL, 資料月份 TEXT NOT NULL,
            {skill_defs}, source_file TEXT, source_row INTEGER, source_payload TEXT,
            FOREIGN KEY(ID, 資料月份) REFERENCES jobs(ID, 資料月份));
        CREATE TABLE IF NOT EXISTS import_runs (
            run_id INTEGER PRIMARY KEY, mode TEXT NOT NULL, source_file TEXT NOT NULL,
            source_rows INTEGER NOT NULL, imported_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS raw_job_sources (
            source_id INTEGER PRIMARY KEY, ID TEXT NOT NULL, 資料月份 TEXT NOT NULL,
            source_file TEXT NOT NULL, source_row INTEGER NOT NULL, source_payload TEXT NOT NULL,
            FOREIGN KEY(ID, 資料月份) REFERENCES jobs(ID, 資料月份));
        CREATE INDEX IF NOT EXISTS idx_skill_snapshot_record ON skill_records(ID, 資料月份, record_id);
        CREATE INDEX IF NOT EXISTS idx_skill_id_snapshot ON skill_records(SKILL_ID, ID, 資料月份);
        CREATE INDEX IF NOT EXISTS idx_skill_category_snapshot ON skill_records(category, ID, 資料月份);
        CREATE INDEX IF NOT EXISTS idx_raw_snapshot ON raw_job_sources(ID, 資料月份);
    """)
    missing = set(JOB_COLUMNS) - columns(conn, "jobs")
    if missing:
        raise ValueError(f"Incompatible jobs schema: missing {sorted(missing)}")
    available = columns(conn, "skill_records")
    required = {"ID", "資料月份", "record_id", "category", "subcategory", "Skill_Name_ZH", "SKILL_ID"}
    if required - available:
        raise ValueError(f"Incompatible skill_records schema: missing {sorted(required - available)}")
    skill_select = [f"sr.{quote(c)}" if c in available else f"NULL AS {quote(c)}"
                    for c in SKILL_COLUMNS if c not in {"category", "subcategory"}]
    conn.execute("DROP VIEW IF EXISTS v_skill_long")
    conn.execute(f"""CREATE VIEW v_skill_long AS SELECT sr.category, sr.subcategory,
        {','.join('j.' + quote(c) for c in JOB_COLUMNS)}, {','.join(skill_select)}, sr.record_id
        FROM jobs j LEFT JOIN skill_records sr
        ON j.ID = sr.ID AND j.資料月份 = sr.資料月份""")


def put_job(conn, row, month=None):
    vals = job_values(row, month)
    existing = conn.execute(f"SELECT {','.join(map(quote, JOB_COLUMNS))} FROM jobs WHERE ID=? AND 資料月份=?",
                            vals[:2]).fetchone()
    if existing is None:
        conn.execute(f"INSERT INTO jobs ({','.join(map(quote, JOB_COLUMNS))}) VALUES ({','.join('?' for _ in vals)})", vals)
    else:
        merged = list(existing)
        for i, v in enumerate(vals[2:], 2):
            if v is not None and existing[i] is not None and v != str(existing[i]):
                raise ValueError(f"Conflicting job snapshot {vals[:2]} in {JOB_COLUMNS[i]}")
            if existing[i] is None:
                merged[i] = v
        if merged != list(existing):
            conn.execute(f"UPDATE jobs SET {','.join(quote(c)+'=?' for c in JOB_COLUMNS[2:])} WHERE ID=? AND 資料月份=?",
                         merged[2:] + vals[:2])
    return vals[:2]


def write_long_rows(conn, rows, source_file="<iterator>", month=None):
    """Every source row becomes one record, including fully blank skill rows and duplicates."""
    count = 0
    cols = ["ID", "資料月份"] + SKILL_COLUMNS + ["source_file", "source_row", "source_payload"]
    sql = f"INSERT INTO skill_records ({','.join(map(quote, cols))}) VALUES ({','.join('?' for _ in cols)})"
    for count, row in enumerate(rows, 1):
        key = put_job(conn, row, month)
        skills = [boolean(value(row, c)) if c == "IS_SOFTWARE" else scalar(value(row, c)) for c in SKILL_COLUMNS]
        conn.execute(sql, key + skills + [source_file, count, payload(row)])
    return count


def iter_source(path, batch_size=5000):
    """No full DataFrame: Parquet batches, CSV iterator, JSONL lines, or read-only XLSX."""
    if batch_size < 1:
        raise ValueError("batch_size must be positive")
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".parquet":
        import pyarrow.parquet as pq
        with pq.ParquetFile(path) as pf:
            for batch in pf.iter_batches(batch_size=batch_size):
                yield from batch.to_pylist()
    elif suffix == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as f:
            yield from csv.DictReader(f)
    elif suffix == ".jsonl":
        with path.open(encoding="utf-8") as f:
            for line in f:
                if line.strip():
                    yield json.loads(line)
    elif suffix == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            rows = wb.active.iter_rows(values_only=True)
            header = next(rows)
            if any(c is None for c in header) or len(set(header)) != len(header):
                raise ValueError("XLSX headers must be nonblank and unique")
            for cells in rows:
                if any(v is not None for v in cells):
                    yield dict(zip(header, cells))
        finally:
            wb.close()
    else:
        raise ValueError("Supported inputs: .parquet, .csv (UTF-8), .jsonl, .xlsx")


def report(conn):
    def n(sql):
        return conn.execute(sql).fetchone()[0]
    result = {
        "job_snapshots": n("SELECT COUNT(*) FROM jobs"),
        "skill_detail_rows": n("SELECT COUNT(*) FROM skill_records"),
        "view_rows": n("SELECT COUNT(*) FROM v_skill_long"),
        "jobs_without_detail_rows": n("""SELECT COUNT(*) FROM jobs j WHERE NOT EXISTS
            (SELECT 1 FROM skill_records sr WHERE sr.ID=j.ID AND sr.資料月份=j.資料月份)"""),
        "distinct_nonempty_skill_ids": n("SELECT COUNT(DISTINCT SKILL_ID) FROM skill_records WHERE TRIM(SKILL_ID)<>''"),
        "distinct_nonempty_skill_names": n("SELECT COUNT(DISTINCT Skill_Name_ZH) FROM skill_records WHERE TRIM(Skill_Name_ZH)<>''"),
        "null_skill_names": n("SELECT COUNT(*) FROM skill_records WHERE Skill_Name_ZH IS NULL"),
        "empty_skill_names": n("SELECT COUNT(*) FROM skill_records WHERE Skill_Name_ZH=''"),
        "raw_source_job_snapshots": n("SELECT COUNT(*) FROM (SELECT DISTINCT ID, 資料月份 FROM raw_job_sources)"),
        "jobs_not_in_raw_sources": n("""SELECT COUNT(*) FROM jobs j WHERE NOT EXISTS
            (SELECT 1 FROM raw_job_sources r WHERE r.ID=j.ID AND r.資料月份=j.資料月份)"""),
    }
    runs = conn.execute("SELECT mode, source_file, source_rows FROM import_runs ORDER BY run_id").fetchall()
    result["imports"] = [dict(zip(["mode", "source_file", "source_rows"], r)) for r in runs]
    result["scope"] = ("raw jobs imported; no-detail jobs are candidates for review, not proven extraction failures"
                       if any(r[0] == "raw-jobs" for r in runs)
                       else "long-only; original jobs absent from the source cannot be counted")
    result["integrity_check"] = conn.execute("PRAGMA integrity_check").fetchone()[0]
    result["foreign_key_errors"] = len(conn.execute("PRAGMA foreign_key_check").fetchall())
    return result


def import_source(source, db_path, *, mode="long", batch_size=5000, month=None, overwrite=False):
    """Build/update a temporary DB; failures leave the existing target unchanged."""
    if mode not in {"long", "raw-jobs"}:
        raise ValueError("Invalid mode")
    source, target = Path(source).resolve(), Path(db_path).resolve()
    if source == target:
        raise ValueError("Input and output must differ")
    if not source.is_file():
        raise FileNotFoundError(source)
    if target.exists() and mode == "long" and not overwrite:
        raise FileExistsError("Target exists; use --overwrite to replace it")
    target.parent.mkdir(parents=True, exist_ok=True)
    fd, temp = tempfile.mkstemp(prefix=target.name + ".", suffix=".tmp", dir=target.parent)
    os.close(fd)
    try:
        with closing(sqlite3.connect(temp)) as conn:
            if mode == "raw-jobs" and target.exists():
                with closing(sqlite3.connect(target.as_uri() + "?mode=ro", uri=True)) as old:
                    old.backup(conn)
            initialize(conn)
            before = conn.execute("SELECT COUNT(*) FROM skill_records").fetchone()[0]
            rows = iter_source(source, batch_size)
            with conn:
                if mode == "long":
                    count = write_long_rows(conn, rows, str(source), month)
                else:
                    count = 0
                    for count, row in enumerate(rows, 1):
                        key = put_job(conn, row, month)
                        conn.execute("""INSERT INTO raw_job_sources
                            (ID, 資料月份, source_file, source_row, source_payload) VALUES (?,?,?,?,?)""",
                                     key + [str(source), count, payload(row)])
                conn.execute("INSERT INTO import_runs(mode,source_file,source_rows,imported_at) VALUES (?,?,?,?)",
                             (mode, str(source), count, datetime.now().isoformat()))
            result = report(conn)
            after = result["skill_detail_rows"]
            if after - before != (count if mode == "long" else 0):
                raise ValueError("Source/detail reconciliation failed")
            if result["integrity_check"] != "ok" or result["foreign_key_errors"]:
                raise ValueError("Database integrity checks failed")
        # Do not replace a live WAL database: its sidecars belong to the old file.
        if any(Path(str(target) + suffix).exists() for suffix in ("-wal", "-shm")):
            raise RuntimeError("Close database viewers/writers before replacing the DB (WAL/SHM exists)")
        os.replace(temp, target)
        result.update(db_path=str(target), source_rows=count, mode=mode)
        return result
    finally:
        if os.path.exists(temp):
            os.unlink(temp)
